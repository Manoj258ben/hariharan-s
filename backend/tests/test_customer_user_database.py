import io

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.database import Base, get_customer_db, get_db, get_user_db
from backend.app.main import app
from backend.app.models import (
    Customer,
    InventoryLot,
    InventoryMovement,
    Product,
    Sale,
    Supplier,
    UserPurchase,
)


@pytest.fixture
def client():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    customer_engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    user_engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    session_factory = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    customer_session_factory = sessionmaker(autocommit=False, autoflush=False, bind=customer_engine)
    user_session_factory = sessionmaker(autocommit=False, autoflush=False, bind=user_engine)
    separate_tables = {"customers", "customer_purchases", "user_purchases"}
    Base.metadata.create_all(
        bind=engine,
        tables=[table for table in Base.metadata.sorted_tables if table.name not in separate_tables],
    )
    Base.metadata.create_all(
        bind=customer_engine,
        tables=[Base.metadata.tables["customers"], Base.metadata.tables["customer_purchases"]],
    )
    Base.metadata.create_all(
        bind=user_engine,
        tables=[Base.metadata.tables["user_purchases"]],
    )

    db = session_factory()
    db.add(Supplier(supplier_id="SUP_TEST", supplier_name="Regression Supplier"))
    db.add(Product(
        product_id="P_TEST",
        barcode="8901234567890",
        product_name="Regression Product",
        category="Test",
        purchase_price=5.0,
        selling_price=10.0,
        current_stock=20.0,
        supplier_id="SUP_TEST",
        storage_capacity_units=100,
    ))
    customer_db = customer_session_factory()
    customer_db.add(Customer(
        customer_id="C_TEST",
        customer_name="Test Customer",
        total_spent=0.0,
        total_purchases_count=0,
    ))
    customer_db.commit()
    customer_db.close()
    user_db = user_session_factory()
    user_db.add(UserPurchase(
        user_purchase_id="UP_SEED",
        user_name="Test Staff",
        product_id="P_TEST",
        product_name="Regression Product",
        supplier_id="SUP_TEST",
        supplier_name="Regression Supplier",
        quantity_bought=1.0,
        purchase_rate=5.0,
        suggested_selling_rate=10.0,
        total_cost=5.0,
        invoice_number="INV-SEED",
        payment_status="PAID",
        purchase_date="2026-10-01",
    ))
    user_db.commit()
    user_db.close()
    db.commit()
    db.close()

    def override_get_db():
        session = session_factory()
        try:
            yield session
        finally:
            session.close()

    def override_get_customer_db():
        session = customer_session_factory()
        try:
            yield session
        finally:
            session.close()

    def override_get_user_db():
        session = user_session_factory()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[get_customer_db] = override_get_customer_db
    app.dependency_overrides[get_user_db] = override_get_user_db
    try:
        yield TestClient(app), session_factory, customer_session_factory, user_session_factory
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_customer_db, None)
        app.dependency_overrides.pop(get_user_db, None)
        Base.metadata.drop_all(bind=engine)
        Base.metadata.drop_all(bind=customer_engine, tables=[
            Base.metadata.tables["customers"],
            Base.metadata.tables["customer_purchases"],
        ])
        Base.metadata.drop_all(bind=user_engine, tables=[Base.metadata.tables["user_purchases"]])
        engine.dispose()
        customer_engine.dispose()
        user_engine.dispose()


def test_customer_barcode_purchase_updates_customer_sales_and_stock(client):
    test_client, session_factory, customer_session_factory, _ = client
    response = test_client.post(
        "/api/customer-user-db/customer-purchases",
        json={
            "customer_id": "C_TEST",
            "barcode": "8901234567890",
            "quantity": 2,
            "payment_method": "Cash",
            "scan_source": "CAMERA_SCAN",
        },
    )

    assert response.status_code == 201
    purchase = response.json()
    assert purchase["customer_name"] == "Test Customer"
    assert purchase["rate"] == 10.0
    assert purchase["total_amount"] == 20.0
    assert purchase["scan_source"] == "CAMERA_SCAN"
    db = session_factory()
    customer_db = customer_session_factory()
    try:
        product = db.query(Product).filter_by(product_id="P_TEST").one()
        customer = customer_db.query(Customer).filter_by(customer_id="C_TEST").one()
        assert product.current_stock == 18.0
        assert customer.total_spent == 20.0
        assert customer.total_purchases_count == 1
        assert db.query(Sale).filter_by(product_id="P_TEST").count() == 1
        movement = db.query(InventoryMovement).filter_by(reference_id=purchase["purchase_id"]).one()
        assert movement.quantity_change == -2.0
    finally:
        db.close()
        customer_db.close()


def test_staff_procurement_updates_stock_rates_and_history(client):
    test_client, session_factory, _, user_session_factory = client
    response = test_client.post(
        "/api/customer-user-db/user-purchases",
        json={
            "user_name": "Test Staff",
            "product_id": "P_TEST",
            "supplier_id": "SUP_TEST",
            "quantity_bought": 4,
            "purchase_rate": 6,
            "suggested_selling_rate": 12,
            "invoice_number": "INV-TEST-2",
            "update_product_rates": True,
        },
    )

    assert response.status_code == 201
    assert response.json()["total_cost"] == 24.0
    db = session_factory()
    user_db = user_session_factory()
    try:
        product = db.query(Product).filter_by(product_id="P_TEST").one()
        assert product.current_stock == 24.0
        assert product.purchase_price == 6.0
        assert product.selling_price == 12.0
        assert user_db.query(UserPurchase).filter_by(invoice_number="INV-TEST-2").count() == 1
        lot = db.query(InventoryLot).filter_by(product_id="P_TEST").one()
        assert lot.quantity_remaining == 4.0
    finally:
        db.close()
        user_db.close()


def test_excel_and_csv_analysis_use_quantity_and_all_workbook_sheets(client):
    test_client, _, _, _ = client
    csv_response = test_client.post(
        "/api/customer-user-db/analyze-excel",
        files={"file": ("products.CSV", b"product_name,category,quantity,unit_price\nTea,Food,2,10\nRice,Food,3,5\n", "text/csv")},
    )
    assert csv_response.status_code == 200
    csv_analysis = csv_response.json()["analysis"]
    assert csv_analysis["total_monetary_value"] == 35.0
    assert csv_analysis["average_rate"] == 7.5
    assert csv_analysis["top_categories"] == [{"category": "Food", "count": 2}]

    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Products"
    first.append(["Product Name", "Quantity", "Unit Price"])
    first.append(["Oil", 2, 10])
    second = workbook.create_sheet("Sales")
    second.append(["Product Name", "Quantity", "Unit Price"])
    second.append(["Flour", 3, 5])
    content = io.BytesIO()
    workbook.save(content)

    xlsx_response = test_client.post(
        "/api/customer-user-db/analyze-excel",
        files={"file": ("inventory.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert xlsx_response.status_code == 200
    xlsx_data = xlsx_response.json()
    assert xlsx_data["total_rows"] == 2
    assert xlsx_data["sheets"] == ["Products", "Sales"]
    assert xlsx_data["analysis"]["total_monetary_value"] == 35.0


def test_reorder_recommendations_include_storage_capacity(client):
    test_client, _, _, _ = client
    response = test_client.get("/api/reorder/recommendations?as_of_date=2026-10-01")

    assert response.status_code == 200
    recommendation = next(item for item in response.json() if item["product_id"] == "P_TEST")
    assert recommendation["storage_capacity_units"] == 100


def test_database_route_aliases_excel_export_and_purchase_order(client):
    test_client, _, _, _ = client
    for prefix in ("/api/database", "/api/customer-user-db"):
        assert test_client.get(f"{prefix}/summary").status_code == 200

    export = test_client.get("/api/customer-user-db/export-excel")
    assert export.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(export.content), read_only=True)
    assert workbook.sheetnames == [
        "Customer Scan Purchases",
        "User Procurement Orders",
        "Products & Active Rates",
        "Rate Modification History",
        "Customers Directory",
    ]
    workbook.close()

    order = test_client.post(
        "/api/purchase-orders",
        json={
            "supplier_id": "SUP_TEST",
            "items": [{"product_id": "P_TEST", "quantity_ordered": 5, "unit_cost": 5}],
        },
    )
    assert order.status_code == 201
    assert order.json()["status"] == "ORDERED"
    assert order.json()["items"][0]["quantity_ordered"] == 5
