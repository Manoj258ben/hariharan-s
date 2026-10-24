"""
StockSense Dynamic Excel & CSV Report Generator
Implements Section 12 of the master prompt generating database-backed reports using OpenPyXL.
"""
import io
import csv
import datetime
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session
from ..models import Product, Sale, InventoryLot, PurchaseOrder, InventoryMovement
from .reorder import get_all_reorder_recommendations

def create_styled_excel_workbook(sheets_data: Dict[str, Dict[str, Any]]) -> bytes:
    """
    Creates a styled multi-sheet Excel workbook with headers, colors, and formatting.
    sheets_data format:
    {
      "SheetName": {
        "headers": ["Col 1", "Col 2", ...],
        "rows": [[val1, val2, ...], ...]
      }
    }
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    data_font = Font(name="Arial", size=9)
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for sheet_name, content in sheets_data.items():
        ws = wb.create_sheet(title=sheet_name[:31])  # Excel 31 char limit
        ws.views.sheetView[0].showGridLines = True

        headers = content["headers"]
        rows = content["rows"]

        # Write header
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Write data rows
        for row in rows:
            ws.append(row)

        for row_idx in range(2, len(rows) + 2):
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = border_thin

        # Auto-fit columns
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_inventory_csv(db: Session) -> str:
    products = db.query(Product).filter(Product.is_archived == False).all()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Product ID", "Barcode", "SKU", "Product Name", "Category", "Brand", "Unit", "Cost Price", "Sell Price", "Current Stock", "Supplier", "Lead Time Days", "Shelf Life Days"])
    for p in products:
        writer.writerow([
            p.product_id, p.barcode or "", p.sku or "", p.product_name, p.category, p.brand,
            p.unit, p.purchase_price, p.selling_price, p.current_stock,
            p.supplier.supplier_name if p.supplier else "", p.supplier_lead_time_days, p.shelf_life_days
        ])
    return output.getvalue()


def export_reorder_csv(db: Session, as_of_date: str = "2026-09-30") -> str:
    recommendations = get_all_reorder_recommendations(db, as_of_date=as_of_date)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Rank", "Product ID", "Product Name", "Category", "Urgency",
        "Daily Demand", "Lead Time (Days)", "Days Stock Left", "Current Stock",
        "Recommended Order Qty", "Pack Size", "Estimated Cost", "Expected Profit", "Supplier"
    ])
    for r in recommendations:
        writer.writerow([
            r["rank"], r["product_id"], r["product_name"], r["category"], r["urgency"],
            r["daily_demand"], r["supplier_lead_time_days"], r["days_of_stock_left"], r["usable_stock"],
            r["suggested_order_quantity"], r["supplier_pack_size"], r["estimated_order_cost"],
            r["expected_gross_profit"], r["supplier_name"]
        ])
    return output.getvalue()


def export_full_excel_report(db: Session, as_of_date: str = "2026-09-30") -> bytes:
    # 1. Inventory Sheet
    products = db.query(Product).filter(Product.is_archived == False).all()
    inv_headers = ["Product ID", "Barcode", "Name", "Category", "Current Stock", "Unit", "Cost (INR)", "Price (INR)", "Lead Time (d)", "Shelf Life (d)"]
    inv_rows = [
        [p.product_id, p.barcode or "", p.product_name, p.category, p.current_stock, p.unit, p.purchase_price, p.selling_price, p.supplier_lead_time_days, p.shelf_life_days]
        for p in products
    ]

    # 2. Reorder Recommendations Sheet
    reorder_list = get_all_reorder_recommendations(db, as_of_date=as_of_date)
    reorder_headers = ["Rank", "Product ID", "Product Name", "Category", "Urgency", "Daily Demand", "Lead Time", "Current Stock", "Reorder Qty", "Est Cost", "Exp Profit", "Supplier"]
    reorder_rows = [
        [r["rank"], r["product_id"], r["product_name"], r["category"], r["urgency"], r["daily_demand"], r["supplier_lead_time_days"], r["usable_stock"], r["suggested_order_quantity"], r["estimated_order_cost"], r["expected_gross_profit"], r["supplier_name"]]
        for r in reorder_list
    ]

    # 3. Expiry Alerts Sheet
    lots = db.query(InventoryLot).all()
    as_of_dt = datetime.datetime.strptime(as_of_date, "%Y-%m-%d").date()
    lot_headers = ["Lot ID", "Product ID", "Expiry Date", "Remaining Qty", "Unit Cost", "Cost at Risk", "Days Left", "Status"]
    lot_rows = []
    for l in lots:
        exp_dt = datetime.datetime.strptime(l.expiry_date, "%Y-%m-%d").date()
        days_left = (exp_dt - as_of_dt).days
        status = "EXPIRED" if days_left <= 0 else ("EXPIRING SOON" if days_left <= 7 else "SAFE")
        cost_at_risk = round(l.quantity_remaining * l.unit_cost, 2)
        lot_rows.append([l.lot_id, l.product_id, l.expiry_date, l.quantity_remaining, l.unit_cost, cost_at_risk, days_left, status])

    # 4. Forecast Accuracy Sheet
    from .forecasting import get_product_daily_sales_series, backtest_product_models
    acc_headers = ["Product ID", "Product Name", "Category", "Evaluated Observations", "Actual Sales Units", "Baseline MAE", "Baseline WAPE (%)", "Winning Model"]
    acc_rows = []
    for p in products:
        series = get_product_daily_sales_series(db, p.product_id, as_of_date=as_of_date, days_back=90)
        bt = backtest_product_models(series, holdout_days=14)
        if bt:
            acc_rows.append([
                p.product_id, p.product_name, p.category, bt["n_observations"],
                sum(bt["actuals"]), bt["baseline_mae"], f"{bt['baseline_wape']}%" if bt['baseline_wape'] is not None else "N/A", bt["winning_model"]
            ])

    sheets_data = {
        "Inventory": {"headers": inv_headers, "rows": inv_rows},
        "Reorder Plan": {"headers": reorder_headers, "rows": reorder_rows},
        "Forecast Accuracy": {"headers": acc_headers, "rows": acc_rows},
        "Expiry Alerts": {"headers": lot_headers, "rows": lot_rows}
    }

    return create_styled_excel_workbook(sheets_data)


def export_forecast_accuracy_excel(db: Session, as_of_date: str = "2026-09-30") -> bytes:
    from .forecasting import get_product_daily_sales_series, backtest_product_models
    products = db.query(Product).filter(Product.is_archived == False).all()
    acc_headers = ["Product ID", "Product Name", "Category", "Evaluated Observations", "Actual Sales Units", "Baseline MAE", "Baseline WAPE (%)", "Winning Model"]
    acc_rows = []
    for p in products:
        series = get_product_daily_sales_series(db, p.product_id, as_of_date=as_of_date, days_back=90)
        bt = backtest_product_models(series, holdout_days=14)
        if bt:
            acc_rows.append([
                p.product_id, p.product_name, p.category, bt["n_observations"],
                sum(bt["actuals"]), bt["baseline_mae"], f"{bt['baseline_wape']}%" if bt['baseline_wape'] is not None else "N/A", bt["winning_model"]
            ])

    sheets_data = {
        "Forecast Accuracy": {"headers": acc_headers, "rows": acc_rows}
    }
    return create_styled_excel_workbook(sheets_data)
